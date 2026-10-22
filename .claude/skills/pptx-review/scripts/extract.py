"""pptx の構造を JSON に書き出す。

使い方: python extract.py <入力.pptx> <出力.json>

スライドごとの図形（スライド番号・shape id・種類・位置/サイズ・塗り・線・効果・
prstGeom と調整値・テキストと書式・揃え・絵文字の有無）、レイアウト、
マスターとテーマ（配色・フォント・背景）を書き出す。
位置とサイズはスライドに対する割合（0〜1）と pt の両方で出す。
"""

import json
import re
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import etree

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}
EMU_PER_PT = 12700

# 絵文字・ピクトグラムとして扱う範囲（記号類の一部を含む）
EMOJI_RE = re.compile(
    "["
    "\U0001F000-\U0001FAFF"
    "\U00002600-\U000027BF"
    "\U00002B00-\U00002BFF"
    "\U0001F900-\U0001F9FF"
    "←-⇿"
    "⌀-⏿"
    "■-◿"
    "️"
    "]"
)


def local(tag):
    return etree.QName(tag).localname


def color_of(el):
    """a:solidFill などの子にある色要素を {"srgb": "RRGGBB"} / {"scheme": "accent1"} にする。"""
    if el is None:
        return None
    for child in el:
        name = local(child.tag)
        if name == "srgbClr":
            out = {"srgb": child.get("val")}
        elif name == "schemeClr":
            out = {"scheme": child.get("val")}
        elif name == "sysClr":
            out = {"sys": child.get("val"), "srgb": child.get("lastClr")}
        elif name == "prstClr":
            out = {"preset": child.get("val")}
        else:
            continue
        mods = {local(m.tag): m.get("val") for m in child}
        if mods:
            out["mods"] = mods
        return out
    return None


def fill_of(parent):
    """spPr / bgPr / rPr などの直下にある塗りを読む。無ければ None（継承）。"""
    if parent is None:
        return None
    for child in parent:
        name = local(child.tag)
        if name == "noFill":
            return {"type": "none"}
        if name == "solidFill":
            return {"type": "solid", "color": color_of(child)}
        if name == "gradFill":
            stops = [
                {"pos": gs.get("pos"), "color": color_of(gs)}
                for gs in child.findall("a:gsLst/a:gs", NS)
            ]
            lin = child.find("a:lin", NS)
            path = child.find("a:path", NS)
            return {
                "type": "gradient",
                "stops": stops,
                "angle": int(lin.get("ang")) / 60000 if lin is not None and lin.get("ang") else None,
                "path": path.get("path") if path is not None else None,
            }
        if name == "blipFill":
            return {"type": "picture"}
        if name == "pattFill":
            return {"type": "pattern", "preset": child.get("prst")}
        if name == "grpFill":
            return {"type": "group"}
    return None


def line_of(sppr):
    ln = sppr.find("a:ln", NS) if sppr is not None else None
    if ln is None:
        return None
    out = {"width_pt": round(int(ln.get("w")) / EMU_PER_PT, 2) if ln.get("w") else None}
    fill = fill_of(ln)
    if fill:
        out["fill"] = fill
    dash = ln.find("a:prstDash", NS)
    if dash is not None:
        out["dash"] = dash.get("val")
    return out


def effects_of(parent):
    if parent is None:
        return None
    lst = parent.find("a:effectLst", NS)
    if lst is None:
        return None
    effects = []
    for e in lst:
        item = {"type": local(e.tag)}
        for attr in ("blurRad", "dist", "rad"):
            if e.get(attr):
                item[attr + "_pt"] = round(int(e.get(attr)) / EMU_PER_PT, 2)
        c = color_of(e)
        if c:
            item["color"] = c
        effects.append(item)
    return effects  # 空リストは「効果を明示的に消している」


def style_refs(sp_el):
    """p:style（テーマの書式を参照する指定）。effectRef の idx>0 はテーマの影などがかかる。"""
    style = sp_el.find("p:style", NS)
    if style is None:
        return None
    out = {}
    for name in ("lnRef", "fillRef", "effectRef", "fontRef"):
        ref = style.find(f"a:{name}", NS)
        if ref is not None:
            out[name] = ref.get("idx")
    return out


def geometry_of(sppr):
    if sppr is None:
        return None
    prst = sppr.find("a:prstGeom", NS)
    if prst is not None:
        adj = {gd.get("name"): gd.get("fmla") for gd in prst.findall("a:avLst/a:gd", NS)}
        return {"preset": prst.get("prst"), "adjust": adj or None}
    if sppr.find("a:custGeom", NS) is not None:
        return {"preset": "custom"}
    return None


def emojis_in(text):
    return sorted(set(EMOJI_RE.findall(text))) if text else []


def text_of(shape):
    if not shape.has_text_frame:
        return None
    tf = shape.text_frame
    body_pr = tf._txBody.find("a:bodyPr", NS)
    paras = []
    for p in tf.paragraphs:
        text = "".join(r.text for r in p.runs) or p.text
        if not text.strip():
            continue
        ppr = p._p.find("a:pPr", NS)
        runs = []
        for r in p.runs:
            rpr = r._r.find("a:rPr", NS)
            runs.append({
                "text": r.text,
                "size_pt": r.font.size.pt if r.font.size else None,
                "bold": r.font.bold,
                "font": r.font.name,
                "fill": fill_of(rpr),
            })
        paras.append({
            "text": text,
            "level": p.level,
            "align": ppr.get("algn") if ppr is not None else None,
            "bullet": (
                None if ppr is None else
                "none" if ppr.find("a:buNone", NS) is not None else
                "char" if ppr.find("a:buChar", NS) is not None else
                "auto" if ppr.find("a:buAutoNum", NS) is not None else None
            ),
            "runs": runs,
            "emoji": emojis_in(text),
        })
    if not paras:
        return None
    return {
        "anchor": body_pr.get("anchor") if body_pr is not None else None,
        "autofit": (
            None if body_pr is None else
            "shrink" if body_pr.find("a:normAutofit", NS) is not None else
            "resize" if body_pr.find("a:spAutoFit", NS) is not None else
            "none" if body_pr.find("a:noAutofit", NS) is not None else None
        ),
        "paragraphs": paras,
    }


def box_of(shape, sw, sh):
    if shape.left is None or shape.width is None:
        return None
    return {
        "x": round(shape.left / sw, 4),
        "y": round(shape.top / sh, 4),
        "w": round(shape.width / sw, 4),
        "h": round(shape.height / sh, 4),
        "x_pt": round(shape.left / EMU_PER_PT, 1),
        "y_pt": round(shape.top / EMU_PER_PT, 1),
        "w_pt": round(shape.width / EMU_PER_PT, 1),
        "h_pt": round(shape.height / EMU_PER_PT, 1),
    }


def shape_of(shape, sw, sh):
    el = shape._element
    sppr = el.find("p:spPr", NS)
    if sppr is None:
        sppr = el.find("p:grpSpPr", NS)
    item = {
        "id": shape.shape_id,
        "name": shape.name,
        "kind": str(shape.shape_type).split(".")[-1].split(" ")[0] if shape.shape_type else local(el.tag),
        "box": box_of(shape, sw, sh),
    }
    if shape.is_placeholder:
        item["placeholder"] = {"type": str(shape.placeholder_format.type).split(".")[-1].split(" ")[0],
                               "idx": shape.placeholder_format.idx}
    optional = {
        "geometry": geometry_of(sppr),
        "fill": fill_of(sppr),
        "line": line_of(sppr),
        "effects": effects_of(sppr),
        "style": style_refs(el),
        "text": text_of(shape),
    }
    item.update({k: v for k, v in optional.items() if v is not None})
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        try:
            item["image"] = {"content_type": shape.image.content_type}
        except Exception:
            item["image"] = {"content_type": None}
    if getattr(shape, "has_table", False) and shape.has_table:
        rows = [[c.text for c in r.cells] for r in shape.table.rows]
        item["table"] = {"rows": len(rows), "cols": len(rows[0]) if rows else 0, "cells": rows}
    if getattr(shape, "has_chart", False) and shape.has_chart:
        item["chart"] = {"type": str(shape.chart.chart_type).split(".")[-1].split(" ")[0]}
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        item["children"] = [shape_of(s, sw, sh) for s in shape.shapes]
    return item


def background_of(cSld_parent):
    bg = cSld_parent.find("p:cSld/p:bg", NS)
    if bg is None:
        return None
    bgpr = bg.find("p:bgPr", NS)
    if bgpr is not None:
        return {"fill": fill_of(bgpr), "effects": effects_of(bgpr)}
    ref = bg.find("p:bgRef", NS)
    if ref is not None:
        return {"theme_ref": ref.get("idx"), "color": color_of(ref)}
    return None


def theme_of(master):
    try:
        part = master.part.part_related_by(RT.THEME)
    except KeyError:
        return None
    root = etree.fromstring(part.blob)
    clr = root.find(".//a:clrScheme", NS)
    colors = {}
    if clr is not None:
        for c in clr:
            colors[local(c.tag)] = color_of(c)
    fonts = {}
    for kind in ("majorFont", "minorFont"):
        f = root.find(f".//a:fontScheme/a:{kind}", NS)
        if f is not None:
            fonts[kind] = {
                "latin": f.find("a:latin", NS).get("typeface") if f.find("a:latin", NS) is not None else None,
                "ea": f.find("a:ea", NS).get("typeface") if f.find("a:ea", NS) is not None else None,
            }
    effect_styles = []
    for es in root.findall(".//a:fmtScheme/a:effectStyleLst/a:effectStyle", NS):
        effect_styles.append([e["type"] for e in (effects_of(es) or [])])
    fill_styles = []
    bg_lst = root.find(".//a:fmtScheme/a:bgFillStyleLst", NS)
    if bg_lst is not None:
        for child in bg_lst:
            wrapper = etree.Element("w")
            wrapper.append(etree.fromstring(etree.tostring(child)))
            fill_styles.append(fill_of(wrapper))
    return {
        "name": root.get("name"),
        "colors": colors,
        "fonts": fonts,
        "effect_styles": effect_styles,
        "bg_fill_styles": fill_styles,
    }


def main(src, dst):
    prs = Presentation(src)
    sw, sh = prs.slide_width, prs.slide_height

    masters = []
    for mi, master in enumerate(prs.slide_masters):
        masters.append({
            "index": mi,
            "background": background_of(master._element),
            "theme": theme_of(master),
            "shapes": [shape_of(s, sw, sh) for s in master.shapes],
            "layouts": [
                {
                    "name": layout.name,
                    "background": background_of(layout._element),
                    "shapes": [shape_of(s, sw, sh) for s in layout.shapes],
                }
                for layout in master.slide_layouts
            ],
        })

    master_index = {m.part.partname: i for i, m in enumerate(prs.slide_masters)}
    slides = []
    for n, slide in enumerate(prs.slides, start=1):
        layout = slide.slide_layout
        slides.append({
            "number": n,
            "layout": layout.name,
            "master": master_index.get(layout.slide_master.part.partname),
            "background": background_of(slide._element),
            "has_notes": slide.has_notes_slide and bool(slide.notes_slide.notes_text_frame.text.strip()),
            "shapes": [shape_of(s, sw, sh) for s in slide.shapes],
        })

    out = {
        "file": src,
        "slide_size_pt": {"w": round(sw / EMU_PER_PT, 1), "h": round(sh / EMU_PER_PT, 1)},
        "slide_count": len(slides),
        "masters": masters,
        "slides": slides,
    }
    with open(dst, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=1)
    print(f"{len(slides)} slides -> {dst}")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
    if len(sys.argv) != 3:
        print("usage: python extract.py <input.pptx> <output.json>", file=sys.stderr)
        sys.exit(2)
    main(sys.argv[1], sys.argv[2])
