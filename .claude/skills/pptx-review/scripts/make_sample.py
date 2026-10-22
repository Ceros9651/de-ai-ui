"""検証用に「AIっぽい」要素を詰め込んだ pptx を作る。

使い方: python make_sample.py <出力.pptx>

含めるもの:
- マスター背景の紫〜青グラデーション（全スライドに効く）
- 絵文字付きの中央寄せタイトル
- 丸いアイコン＋見出し＋短文の角丸カード3枚（影付き）
- AI文体の箇条書き（全部中央寄せ）
- グラデーション塗りの統計カード4枚
- 発表者ノート（変更されないことの確認用）
"""

import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

GRADIENT_BG = f"""
<p:bg xmlns:p="{P}" xmlns:a="{A}">
  <p:bgPr>
    <a:gradFill rotWithShape="1">
      <a:gsLst>
        <a:gs pos="0"><a:srgbClr val="6D28D9"/></a:gs>
        <a:gs pos="100000"><a:srgbClr val="2563EB"/></a:gs>
      </a:gsLst>
      <a:lin ang="2700000" scaled="0"/>
    </a:gradFill>
    <a:effectLst/>
  </p:bgPr>
</p:bg>
"""

SHADOW = f"""
<a:effectLst xmlns:a="{A}">
  <a:outerShdw blurRad="190500" dist="63500" dir="5400000" algn="t" rotWithShape="0">
    <a:srgbClr val="000000"><a:alpha val="35000"/></a:srgbClr>
  </a:outerShdw>
</a:effectLst>
"""


def gradient_fill(c1, c2):
    return etree.fromstring(f"""
<a:gradFill xmlns:a="{A}" rotWithShape="1">
  <a:gsLst>
    <a:gs pos="0"><a:srgbClr val="{c1}"/></a:gs>
    <a:gs pos="100000"><a:srgbClr val="{c2}"/></a:gs>
  </a:gsLst>
  <a:lin ang="5400000" scaled="0"/>
</a:gradFill>""")


def add_shadow(shape):
    sppr = shape._element.spPr
    old = sppr.find(qn("a:effectLst"))
    if old is not None:
        sppr.remove(old)
    sppr.append(etree.fromstring(SHADOW))


def set_text(shape, lines, size, bold=False, align=PP_ALIGN.CENTER, color=WHITE):
    tf = shape.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color


def textbox(slide, x, y, w, h, lines, size, **kw):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    set_text(tb, lines, size, **kw)
    return tb


def card(slide, x, y, w, h, fill="FFFFFF"):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = 0.12
    shp.fill.solid()
    shp.fill.fore_color.rgb = RGBColor.from_string(fill)
    shp.line.fill.background()
    add_shadow(shp)
    return shp


def main(dst):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    master = prs.slide_masters[0]
    csld = master._element.find(qn("p:cSld"))
    old_bg = csld.find(qn("p:bg"))
    if old_bg is not None:
        csld.remove(old_bg)  # p:bg が2つあると PowerPoint で開けなくなる
    csld.insert(0, etree.fromstring(GRADIENT_BG))

    blank = prs.slide_layouts[6]

    # 1. 表紙
    s = prs.slides.add_slide(blank)
    textbox(s, 1, 2.4, 11.333, 1.4, ["🚀 次世代AIプラットフォームで業務を革新"], 44, bold=True)
    textbox(s, 1, 3.9, 11.333, 0.8, ["✨ シームレスな体験を、すべての人に ✨"], 22)
    s.notes_slide.notes_text_frame.text = "ここは自分のメモ。変更されないこと。"

    # 2. アイコン付き角丸カード3枚
    s = prs.slides.add_slide(blank)
    textbox(s, 1, 0.6, 11.333, 1, ["💡 3つの主要な価値"], 36, bold=True)
    items = [
        ("⚡", "圧倒的なスピード", "AIが業務プロセスを最適化し、効率を大幅に向上させます。"),
        ("🔒", "強固なセキュリティ", "エンタープライズグレードの安全性で、安心してご利用いただけます。"),
        ("🌐", "シームレスな連携", "既存のツールとシームレスに統合し、革新的な体験を実現します。"),
    ]
    for i, (icon, head, body) in enumerate(items):
        x = 1 + i * 3.9
        card(s, x, 2.0, 3.5, 4.3)
        circle = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 1.25), Inches(2.4), Inches(1), Inches(1))
        circle.fill.solid()
        circle.fill.fore_color.rgb = RGBColor(0xED, 0xE9, 0xFE)
        circle.line.fill.background()
        set_text(circle, [icon], 28, color=RGBColor(0x6D, 0x28, 0xD9))
        textbox(s, x + 0.2, 3.6, 3.1, 0.6, [head], 22, bold=True, color=RGBColor(0x1F, 0x29, 0x37))
        textbox(s, x + 0.2, 4.3, 3.1, 1.6, [body], 15, color=RGBColor(0x4B, 0x55, 0x63))

    # 3. AI文体の箇条書き（中央寄せ）
    s = prs.slides.add_slide(blank)
    textbox(s, 1, 0.6, 11.333, 1, ["📊 導入のメリット"], 36, bold=True)
    textbox(s, 1.5, 1.9, 10.333, 4.8, [
        "• 革新的なソリューションで業務効率を最大化",
        "• シームレスな統合により、さらなる生産性向上を実現",
        "• 重要なのは、データドリブンな意思決定の加速",
        "• 包括的なサポート体制で、安心の導入をサポート",
        "• 次世代のワークフローを、今ここに",
    ], 22)

    # 4. グラデーションの統計カード4枚
    s = prs.slides.add_slide(blank)
    textbox(s, 1, 0.6, 11.333, 1, ["🎯 導入実績"], 36, bold=True)
    stats = [("98%", "顧客満足度"), ("3倍", "作業スピード"), ("500社+", "導入企業"), ("24/7", "サポート")]
    for i, (num, label) in enumerate(stats):
        x = 1 + i * 2.9
        c = card(s, x, 2.4, 2.6, 2.6)
        sppr = c._element.spPr
        sppr.remove(sppr.find(qn("a:solidFill")))
        sppr.insert(2, gradient_fill("8B5CF6", "3B82F6"))
        textbox(s, x, 2.9, 2.6, 1, [num], 40, bold=True)
        textbox(s, x, 3.9, 2.6, 0.6, [label], 16)

    prs.save(dst)
    print(f"sample -> {dst}")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
    if len(sys.argv) != 2:
        print("usage: python make_sample.py <output.pptx>", file=sys.stderr)
        sys.exit(2)
    main(sys.argv[1])
